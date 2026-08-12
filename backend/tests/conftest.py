"""Shared fixtures for the B1 ingestion/parsing/pipeline tests.

Problem-file fixtures are BUILT here, deterministically, into
tests/fixtures/ -- handcrafted so these tests never depend on the parallel
corpusgen track's problem_files renderer timing (docs/plan.md §14 runs the
tracks concurrently). Each is the smallest possible instance of one plan §8
problem kind.

DB-touching tests follow UC2's convention (see tests/test_models_smoke.py):
they run against the real docker-compose Postgres on :5434 with migrations
applied. Since migration 002, documents.sha256 is UNIQUE per run (not
globally), so pytest ingests can never shadow production rows at the DB
level; the remaining hygiene concerns are (a) fixture bytes must never
EQUAL corpusgen output (docs/plan.md §14b R2 -- hence _FIXTURE_MARKER
below) and (b) `free_sha_collisions` cleans matching rows left by earlier
test sessions while explicitly refusing to touch rows belonging to runs
over the repo's real data/ corpora.
"""

from __future__ import annotations

import io
from pathlib import Path

import pikepdf
import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.db.models import Document, ProcessingRun

FIXTURES_DIR = Path(__file__).parent / "fixtures"
_REPO_DATA_DIR = Path(__file__).resolve().parents[2] / "data"

# Appended to every non-empty problem fixture (docs/plan.md §14b R2): a
# trailing byte no corpusgen renderer ever emits, so fixture bytes can never
# sha256-collide with production corpus files. Trailing junk is harmless to
# every consumer here -- zip readers locate the end-of-central-directory by
# scanning backwards, pikepdf tolerates bytes after %%EOF, and the junk
# files are junk either way. The one file that CANNOT carry it is the
# zero-byte fixture (appending anything would un-zero it); it is instead
# protected by per-run dedup (migration 002) plus the data/-guard in
# free_sha_collisions below.
_FIXTURE_MARKER = b"\xa5"


def build_password_pdf_bytes() -> bytes:
    pdf = pikepdf.new()
    pdf.add_blank_page()
    buffer = io.BytesIO()
    pdf.save(buffer, encryption=pikepdf.Encryption(owner="owner-pass", user="user-pass", R=6))
    return buffer.getvalue()


def build_docx_bytes(lines: list[str]) -> bytes:
    doc = DocxDocument()
    for line in lines:
        doc.add_paragraph(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_xlsx_bytes(headers: list[str], rows: list[list[str]], sheet: str = "Sheet1") -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_eml_bytes(
    *, body: str, attachment_name: str, attachment_content: bytes, html_body: str | None = None
) -> bytes:
    from email.message import EmailMessage

    message = EmailMessage()
    message["From"] = "Dana Whitfield <dana.whitfield@example.org>"
    message["To"] = "hr-intake@example.org"
    message["Subject"] = "Employee records attached"
    message["Date"] = "Mon, 15 Jan 2026 12:00:00 -0000"
    message.set_content(body)
    if html_body is not None:
        message.add_alternative(html_body, subtype="html")
    message.add_attachment(
        attachment_content,
        maintype="text",
        subtype="csv",
        filename=attachment_name,
    )
    return message.as_bytes()


# The six handcrafted problem kinds (plan §8's problem_files, minus the
# scan-specific ones): filename -> byte builder. actually_xlsx.pdf is the
# RECOVERABLE wrong-extension trap (routes to the xlsx parser, no
# quarantine); binary_junk.docx is the unrecoverable one (extension claims
# docx, bytes disprove it).
def _problem_file_builders() -> dict[str, bytes]:
    return {
        "password_protected.pdf": build_password_pdf_bytes() + _FIXTURE_MARKER,
        "truncated.docx": build_docx_bytes(["This document will be cut off mid-zip."])[:400]
        + _FIXTURE_MARKER,
        "zero_byte.txt": b"",  # cannot carry the marker; see _FIXTURE_MARKER note
        "actually_xlsx.pdf": build_xlsx_bytes(["Name", "SSN"], [["Casey Vance", "531-24-8817"]])
        + _FIXTURE_MARKER,
        # Distinct byte patterns: identical content would sha256-dedup into
        # ONE documents row and break the per-file quarantine assertions.
        "binary_junk.docx": b"\x00\xffJUNKDATA" * 64 + _FIXTURE_MARKER,
        "noise.bin": b"\x01\xfeNOISEBIN" * 64 + _FIXTURE_MARKER,
    }


@pytest.fixture(scope="session")
def problem_files() -> dict[str, Path]:
    """Materializes the problem fixtures under tests/fixtures/ and returns
    filename -> path. Rebuilt every session (cheap, and the encrypted PDF's
    salts make its bytes fresh each time anyway)."""
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    for filename, content in _problem_file_builders().items():
        path = FIXTURES_DIR / filename
        path.write_bytes(content)
        paths[filename] = path
    return paths


def free_sha_collisions(db: Session, contents: list[bytes]) -> None:
    """Delete documents rows (cascading passages/quarantines) whose sha256
    matches any of `contents`, EXCEPT rows belonging to runs over the repo's
    real data/ corpora. Since migration 002 (sha unique per run) this is
    pure hygiene -- a prior pytest session's leftovers can no longer shadow
    a new ingest -- but the guard is load-bearing: the zero-byte fixture
    necessarily hashes identically to corpusgen's zero-byte problem file,
    and the unguarded version of this function deleted the production run's
    quarantined document (the docs/plan.md §14b R2 incident)."""
    from app.services.ingestion.inventory import sha256_of

    shas = [sha256_of(content) for content in contents]
    protected_prefix = f"{_REPO_DATA_DIR}%"
    doomed = select(Document.id).join(
        ProcessingRun, Document.run_id == ProcessingRun.id
    ).where(
        Document.sha256.in_(shas),
        ProcessingRun.config_snapshot["corpus_path"].astext.not_like(protected_prefix),
    )
    db.execute(delete(Document).where(Document.id.in_(doomed)))
    db.commit()
